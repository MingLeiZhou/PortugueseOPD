#!/usr/bin/env python3
"""Export the canonical Scientific Data LaTeX manuscript to editable DOCX."""

from __future__ import annotations

import argparse
import re
import subprocess
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


PAPER = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = PAPER / "SimPT60_submission_en.docx"
TITLE = "SimPT60: A Traceable Time-Series Power-Flow Dataset for Portugal's High-Voltage Grid"
FONT = "Times New Roman"
CONTENT_WIDTH_DXA = 9360


def set_run_font(run, name: str, size: float, *, bold: bool | None = None) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(key), name)


def configure_paragraph_style(
    document: Document,
    name: str,
    *,
    size: float,
    bold: bool = False,
    italic: bool = False,
    before: float = 0,
    after: float = 0,
    line_spacing: float = 1.0,
    alignment=None,
    keep_with_next: bool = False,
) -> None:
    style = next(style for style in document.styles if style.name == name)
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(key), FONT)
    style.font.color.rgb = RGBColor(0, 0, 0)
    fmt = style.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line_spacing
    fmt.keep_with_next = keep_with_next
    fmt.widow_control = True
    if alignment is not None:
        fmt.alignment = alignment


def ensure_child(parent, tag: str):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def set_width(parent, tag: str, width_dxa: int) -> None:
    width = ensure_child(parent, tag)
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(width_dxa))


def set_border(parent, tag: str, *, value: str, size: int = 0) -> None:
    border = ensure_child(parent, tag)
    border.set(qn("w:val"), value)
    border.set(qn("w:sz"), str(size))
    border.set(qn("w:space"), "0")
    border.set(qn("w:color"), "000000")


def apply_table_geometry(table, widths_dxa: list[int]) -> None:
    """Synchronize Word table, grid, and cell widths for stable rendering."""

    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table widths must sum to {CONTENT_WIDTH_DXA}: {widths_dxa}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    set_width(tbl_pr, "w:tblW", CONTENT_WIDTH_DXA)
    table_indent = ensure_child(tbl_pr, "w:tblInd")
    table_indent.set(qn("w:type"), "dxa")
    table_indent.set(qn("w:w"), "120")
    layout = ensure_child(tbl_pr, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table_borders = ensure_child(tbl_pr, "w:tblBorders")
    for side in ("top", "bottom", "left", "right", "insideH", "insideV"):
        set_border(table_borders, f"w:{side}", value="nil")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for column, width in zip(table.columns, widths_dxa):
        column.width = Twips(width)
    for row in table.rows:
        row.height = None
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row.cells, widths_dxa):
            cell.width = Twips(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            set_width(tc_pr, "w:tcW", width)
            cell_borders = ensure_child(tc_pr, "w:tcBorders")
            for side in ("top", "bottom", "left", "right"):
                set_border(cell_borders, f"w:{side}", value="nil")
            tc_mar = ensure_child(tc_pr, "w:tcMar")
            for side, margin in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
                item = ensure_child(tc_mar, f"w:{side}")
                item.set(qn("w:w"), str(margin))
                item.set(qn("w:type"), "dxa")

    header_pr = table.rows[0]._tr.get_or_add_trPr()
    if header_pr.find(qn("w:tblHeader")) is None:
        header_pr.append(OxmlElement("w:tblHeader"))
    for cell in table.rows[0].cells:
        cell_borders = ensure_child(cell._tc.get_or_add_tcPr(), "w:tcBorders")
        set_border(cell_borders, "w:top", value="single", size=8)
        set_border(cell_borders, "w:bottom", value="single", size=6)
    for cell in table.rows[-1].cells:
        cell_borders = ensure_child(cell._tc.get_or_add_tcPr(), "w:tcBorders")
        set_border(cell_borders, "w:bottom", value="single", size=8)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, FONT, 9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, text, end):
        run._r.append(element)


def prepend_caption_number(paragraph, label: str, number: int, bookmark_id: int) -> None:
    """Prepend a preview-safe caption number and bookmark."""

    if paragraph.text.startswith(f"{label} "):
        return

    label_run = OxmlElement("w:r")
    label_properties = OxmlElement("w:rPr")
    label_properties.append(OxmlElement("w:b"))
    label_run.append(label_properties)
    label_text = OxmlElement("w:t")
    label_text.set(qn("xml:space"), "preserve")
    label_text.text = f"{label} "
    label_run.append(label_text)

    bookmark_start = OxmlElement("w:bookmarkStart")
    bookmark_start.set(qn("w:id"), str(bookmark_id))
    bookmark_start.set(qn("w:name"), f"{'fig' if label == 'Figure' else 'tbl'}{number}")

    number_run = OxmlElement("w:r")
    number_properties = OxmlElement("w:rPr")
    number_properties.append(OxmlElement("w:b"))
    number_run.append(number_properties)
    result = OxmlElement("w:t")
    result.text = str(number)
    number_run.append(result)

    bookmark_end = OxmlElement("w:bookmarkEnd")
    bookmark_end.set(qn("w:id"), str(bookmark_id))

    punctuation_run = OxmlElement("w:r")
    punctuation_properties = OxmlElement("w:rPr")
    punctuation_properties.append(OxmlElement("w:b"))
    punctuation_run.append(punctuation_properties)
    punctuation = OxmlElement("w:t")
    punctuation.set(qn("xml:space"), "preserve")
    punctuation.text = ". "
    punctuation_run.append(punctuation)

    paragraph_xml = paragraph._p
    insertion_index = 1 if paragraph_xml.pPr is not None else 0
    for element in (label_run, bookmark_start, number_run, bookmark_end, punctuation_run):
        paragraph_xml.insert(insertion_index, element)
        insertion_index += 1


def number_display_equations(document: Document) -> int:
    """Add right-aligned numbers to Pandoc's display-math paragraphs."""

    equation_number = 0
    for paragraph in document.paragraphs:
        math_paragraph = paragraph._p.find(qn("m:oMathPara"))
        if math_paragraph is None:
            continue
        math_object = math_paragraph.find(qn("m:oMath"))
        if math_object is None:
            continue

        equation_number += 1
        math_paragraph.remove(math_object)
        paragraph._p.remove(math_paragraph)

        paragraph_properties = paragraph._p.get_or_add_pPr()
        justification = paragraph_properties.find(qn("w:jc"))
        if justification is not None:
            paragraph_properties.remove(justification)
        tabs = paragraph_properties.find(qn("w:tabs"))
        if tabs is not None:
            paragraph_properties.remove(tabs)
        tabs = OxmlElement("w:tabs")
        center_tab = OxmlElement("w:tab")
        center_tab.set(qn("w:val"), "center")
        center_tab.set(qn("w:pos"), str(CONTENT_WIDTH_DXA // 2))
        right_tab = OxmlElement("w:tab")
        right_tab.set(qn("w:val"), "right")
        right_tab.set(qn("w:pos"), str(CONTENT_WIDTH_DXA))
        tabs.extend((center_tab, right_tab))
        paragraph_properties.append(tabs)

        leading_tab = OxmlElement("w:r")
        leading_tab.append(OxmlElement("w:tab"))
        trailing_run = OxmlElement("w:r")
        trailing_properties = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            fonts.set(qn(key), FONT)
        trailing_properties.append(fonts)
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), "21")
        trailing_properties.append(size)
        trailing_run.append(trailing_properties)
        trailing_run.append(OxmlElement("w:tab"))
        number_text = OxmlElement("w:t")
        number_text.text = f"({equation_number})"
        trailing_run.append(number_text)

        paragraph._p.append(leading_tab)
        paragraph._p.append(math_object)
        paragraph._p.append(trailing_run)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.keep_together = True

    return equation_number


def format_docx(path: Path) -> None:
    """Apply a restrained academic Word style and stable layout geometry."""

    document = Document(path)
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)

    settings = document.settings.element
    view = ensure_child(settings, "w:view")
    view.set(qn("w:val"), "print")
    hidden_boundaries = settings.find(qn("w:doNotDisplayPageBoundaries"))
    if hidden_boundaries is not None:
        settings.remove(hidden_boundaries)
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    # Narrative-proposal preset with named academic-manuscript overrides:
    # Times New Roman, black hierarchy, compact 1.12-line body, journal captions.
    for style_name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        configure_paragraph_style(
            document,
            style_name,
            size=10.5,
            after=5,
            line_spacing=1.12,
            alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        )
    configure_paragraph_style(
        document,
        "Title",
        size=18,
        bold=True,
        after=10,
        line_spacing=1.0,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        keep_with_next=True,
    )
    configure_paragraph_style(
        document,
        "Author",
        size=11,
        after=3,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        keep_with_next=True,
    )
    configure_paragraph_style(
        document,
        "Date",
        size=10,
        after=9,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        keep_with_next=True,
    )
    configure_paragraph_style(
        document,
        "Abstract Title",
        size=11,
        bold=True,
        before=3,
        after=4,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        keep_with_next=True,
    )
    configure_paragraph_style(
        document,
        "Abstract",
        size=10,
        after=7,
        line_spacing=1.08,
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )
    configure_paragraph_style(
        document,
        "Heading 1",
        size=14,
        bold=True,
        before=14,
        after=6,
        keep_with_next=True,
    )
    configure_paragraph_style(
        document,
        "Heading 2",
        size=12,
        bold=True,
        before=10,
        after=4,
        keep_with_next=True,
    )
    configure_paragraph_style(
        document,
        "Heading 3",
        size=11,
        bold=True,
        italic=True,
        before=8,
        after=3,
        keep_with_next=True,
    )
    for style_name in ("Caption", "Table Caption"):
        configure_paragraph_style(
            document,
            style_name,
            size=9,
            italic=True,
            before=6,
            after=4,
            line_spacing=1.0,
            alignment=WD_ALIGN_PARAGRAPH.LEFT,
            keep_with_next=True,
        )
    configure_paragraph_style(
        document,
        "Image Caption",
        size=9,
        italic=True,
        before=3,
        after=8,
        line_spacing=1.0,
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )
    configure_paragraph_style(
        document,
        "Captioned Figure",
        size=9,
        after=0,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        keep_with_next=True,
    )
    configure_paragraph_style(
        document,
        "Bibliography",
        size=8.0,
        after=1,
        line_spacing=0.95,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )
    bibliography = document.styles["Bibliography"].paragraph_format
    bibliography.left_indent = Inches(0.25)
    bibliography.first_line_indent = Inches(-0.25)
    configure_paragraph_style(document, "Source Code", size=8.5, after=3, line_spacing=1.0)
    document.styles["Source Code"].font.name = "Courier New"

    title_style_sizes = {
        "Title": 18,
        "Author": 11,
        "Date": 10,
        "Abstract Title": 11,
        "Abstract": 10,
        "Heading 1": 14,
        "Heading 2": 12,
        "Heading 3": 11,
        "Caption": 9,
        "Table Caption": 9,
        "Image Caption": 9,
        "Bibliography": 8,
        "Source Code": 8.5,
    }
    in_references = False
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == "References":
            in_references = True
            continue
        if in_references:
            paragraph.style = document.styles["Bibliography"]
        if re.match(r"^Table \d+—", paragraph.text):
            paragraph.style = document.styles["Table Caption"]
        elif re.match(r"^Figure \d+—", paragraph.text):
            paragraph.style = document.styles["Image Caption"]

    table_number = 0
    figure_number = 0
    bookmark_id = 100
    for paragraph in document.paragraphs:
        if paragraph.style.name == "Table Caption":
            table_number += 1
            prepend_caption_number(paragraph, "Table", table_number, bookmark_id)
            bookmark_id += 1
        elif paragraph.style.name == "Image Caption":
            figure_number += 1
            prepend_caption_number(paragraph, "Figure", figure_number, bookmark_id)
            bookmark_id += 1

    for paragraph in document.paragraphs:
        style_name = paragraph.style.name
        if style_name == "Heading 4":
            # LaTeX \paragraph headings are semantic third-level headings in
            # this manuscript and are unnumbered in the source. Pandoc's global
            # numbering otherwise creates misleading labels such as 3.3.0.1.
            cleaned_heading = re.sub(r"^\d+(?:\.\d+)*\s+", "", paragraph.text)
            paragraph.clear()
            paragraph.add_run(cleaned_heading)
            paragraph.style = document.styles["Heading 3"]
            style_name = "Heading 3"
        size = title_style_sizes.get(style_name, 10.5)
        font_name = "Courier New" if style_name == "Source Code" else FONT
        for run in paragraph.runs:
            set_run_font(run, font_name, size)
        paragraph.paragraph_format.widow_control = True
        if style_name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True
        elif style_name == "Captioned Figure":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
        elif style_name == "Image Caption":
            paragraph.paragraph_format.keep_together = True
        elif style_name == "Table Caption":
            paragraph.paragraph_format.keep_with_next = True
        if style_name in {"Normal", "Body Text", "First Paragraph", "Compact"}:
            contains_unbreakable_content = (
                "_" in paragraph.text
                or "http://" in paragraph.text
                or "https://" in paragraph.text
                or re.search(r"\b[0-9a-fA-F]{32,}\b", paragraph.text) is not None
            )
            if contains_unbreakable_content:
                # Full justification creates excessive gaps around URLs,
                # digests, and file paths. Keep those paragraphs left aligned
                # while ordinary manuscript prose remains justified.
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    number_display_equations(document)

    for table in document.tables:
        table.autofit = True
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table_properties = table._tbl.tblPr
        table_width = ensure_child(table_properties, "w:tblW")
        table_width.set(qn("w:type"), "pct")
        table_width.set(qn("w:w"), "5000")
        table_layout = ensure_child(table_properties, "w:tblLayout")
        table_layout.set(qn("w:type"), "autofit")
        header_properties = table.rows[0]._tr.get_or_add_trPr()
        if header_properties.find(qn("w:tblHeader")) is None:
            header_properties.append(OxmlElement("w:tblHeader"))
        for row_index, row in enumerate(table.rows):
            row.height = None
            row_properties = row._tr.get_or_add_trPr()
            if row_properties.find(qn("w:cantSplit")) is None:
                row_properties.append(OxmlElement("w:cantSplit"))
            for column_index, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cell_properties = cell._tc.get_or_add_tcPr()
                cell_width = cell_properties.find(qn("w:tcW"))
                if cell_width is not None:
                    cell_properties.remove(cell_width)
                cell_margins = ensure_child(cell_properties, "w:tcMar")
                for side, margin in (("top", 70), ("bottom", 70), ("start", 90), ("end", 90)):
                    item = ensure_child(cell_margins, f"w:{side}")
                    item.set(qn("w:w"), str(margin))
                    item.set(qn("w:type"), "dxa")
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.line_spacing = 1.0
                    paragraph.paragraph_format.widow_control = True
                    if column_index == 1 and len(row.cells) >= 3:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    else:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for run in paragraph.runs:
                        set_run_font(run, FONT, 8.0, bold=(row_index == 0))

    image_captions = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style.name == "Image Caption"
    ]
    for index, shape in enumerate(document.inline_shapes):
        aspect_ratio = shape.height / shape.width
        shape.width = Inches(6.15)
        shape.height = int(shape.width * aspect_ratio)
        doc_pr = shape._inline.docPr
        doc_pr.set("descr", image_captions[index] if index < len(image_captions) else f"Figure {index + 1}")
        doc_pr.set("title", f"Figure {index + 1}")

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.clear()
    add_page_number(footer_paragraph)

    document.core_properties.title = TITLE
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.comments = "Editable submission manuscript export"
    document.save(path)


def remove_empty_comment_part(path: Path) -> None:
    """Remove Pandoc's unused comments part from an otherwise clean DOCX."""

    with zipfile.ZipFile(path, "r") as source:
        document_xml = source.read("word/document.xml")
        if b"w:commentReference" in document_xml:
            return

        relationships = source.read("word/_rels/document.xml.rels").decode("utf-8")
        content_types = source.read("[Content_Types].xml").decode("utf-8")
        relationships = re.sub(
            r'<Relationship\b[^>]*Type="http://schemas\.openxmlformats\.org/'
            r'officeDocument/2006/relationships/comments"[^>]*/>',
            "",
            relationships,
        )
        content_types = re.sub(
            r'<Override\b[^>]*PartName="/word/comments\.xml"[^>]*/>',
            "",
            content_types,
        )

        temporary_path = path.with_suffix(".without-comments.docx")
        with zipfile.ZipFile(temporary_path, "w", zipfile.ZIP_DEFLATED) as target:
            for item in source.infolist():
                if item.filename == "word/comments.xml":
                    continue
                if item.filename == "word/_rels/document.xml.rels":
                    target.writestr(item, relationships.encode("utf-8"))
                elif item.filename == "[Content_Types].xml":
                    target.writestr(item, content_types.encode("utf-8"))
                else:
                    target.writestr(item, source.read(item.filename))
    temporary_path.replace(path)


def escape_texttt(value: str) -> str:
    return value.replace("_", r"\_").replace("#", r"\#")


def prepare_latex(source: str) -> str:
    """Replace PDF-only display helpers with constructs Pandoc preserves."""

    source = re.sub(r"\\begin\{thebibliography\}.*?\\end\{thebibliography\}", "", source, flags=re.DOTALL)
    source = re.sub(
        r"\\pandocbounded\{\\includegraphics(?:\[[^\]]*\])?\{([^{}]+)\}\}",
        r"\\includegraphics{\1}",
        source,
    )
    for macro in ("path", "hashvalue"):
        pattern = re.compile(rf"\\{macro}\{{([^{{}}]+)\}}")
        source = pattern.sub(lambda match: rf"\texttt{{{escape_texttt(match.group(1))}}}", source)
    return source


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=PAPER / "main_scidata.tex")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    source_path = args.source.resolve()
    output_path = args.output.resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)

    prepared = prepare_latex(source_path.read_text(encoding="utf-8"))
    with tempfile.NamedTemporaryFile(
        mode="w", suffix=".tex", prefix="pt60_docx_", dir=PAPER, encoding="utf-8", delete=False
    ) as handle:
        handle.write(prepared)
        temporary_source = Path(handle.name)

    command = [
        "pandoc",
        temporary_source.name,
        "--from=latex",
        "--to=docx",
        "--citeproc",
        "--bibliography=references.bib",
        "--csl=../nature.csl",
        "--resource-path=.:figures_final:figures/generated:generated_tables",
        "--lua-filter=scripts/docx_image_filter.lua",
        f"--metadata=title:{TITLE}",
        "--metadata=date:September 2026",
        "--metadata=reference-section-title:References",
        f"--output={output_path}",
    ]
    try:
        subprocess.run(command, cwd=PAPER, check=True)
    finally:
        temporary_source.unlink(missing_ok=True)

    format_docx(output_path)
    remove_empty_comment_part(output_path)

    with zipfile.ZipFile(output_path) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
        document_xml = archive.read("word/document.xml")
    table_count = document_xml.count(b"<w:tbl>")
    print(f"Wrote {output_path}")
    print(f"Embedded media: {len(media)} ({', '.join(Path(name).suffix for name in media)})")
    print(f"Word tables: {table_count}")


if __name__ == "__main__":
    main()
